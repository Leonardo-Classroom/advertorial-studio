"""Turn an uploaded Brief / Proposal deck into structured facts.

Step 1 pulls every scrap of text out of the .pptx (shapes, grouped shapes,
tables, speaker notes) — proposal decks put the campaign's real substance in
tables and notes as often as in title placeholders.

Step 2 asks the model to normalise that into the fact schema. The result is
explicitly *editable* by the operator before generation: it becomes the only
sanctioned source of brand names, product codes and KOL names, so a mistake
here would propagate into every draft.

`extract_images` is the optional third path: the same shape walk, collecting
pictures instead of text so a draft can place real deck imagery rather than a
`（Photo from …）` placeholder the operator has to fill in by hand.
"""
from __future__ import annotations

from pathlib import Path

from core import llm, timeouts

INSTRUCTIONS = """你是行銷簡報的資料整理員。你的工作是從媒體簡報（Media Brief / Proposal）中，
抽出後續撰寫廣編稿所需的「事實」。

最重要的一件事：**分清楚哪些是能寫給消費者看的，哪些是代理商的內部規劃。**
媒體簡報有一大半在講投放計畫——預算、版位、檔期分期（Pre-heat／Launch／Sustain）、
KPI、聲量目標、捷運燈箱點位、導購成效——這些**絕對不能出現在廣編稿裡**，
讀者不需要也不應該知道。它們一律歸到 internal_only，不要混進其他欄位。

寫進 key_message 的必須是「對讀者有意義的話」，不是「對品牌有意義的目標」。
例如簡報寫「強化薄底鞋商品心佔、帶動周慶業績」，那是內部目標；
對讀者有意義的版本是「這雙鞋的薄底設計讓腿看起來更長」。
如果簡報通篇只有內部目標、沒有任何可以直接講給消費者聽的賣點，
key_message 就留空陣列，不要硬把目標改寫成賣點。

其餘鐵則：只抽取簡報中真實出現的內容，絕對不要推測、補充或發明。
簡報沒寫的欄位就留空字串或空陣列。全程使用繁體中文。
只輸出 JSON，不要有任何其他文字或 ``` 標記。"""

TASK = """請從以下簡報文字中抽取結構化事實，輸出這個 JSON 結構：

{{
  "brand": "品牌名稱",
  "campaign": "檔期／活動名稱（僅供內部辨識，不會寫進稿子）",
  "product": ["產品名稱或型號"],
  "product_details": ["產品的具體特徵：配色名稱、材質、鞋型、設計細節、穿起來的感受。這是廣編稿最需要的東西"],
  "slogan": "主標語（簡報中明確寫出的）",
  "angle": "簡報若已經替這篇稿決定了切角、方向或直接指定了標題，照抄那句話。提案簡報常針對每個媒體各設一個切角，若有就抽出來。**簡報沒寫就留空字串，絕對不要自己想一句**",
  "key_message": ["對讀者有意義的賣點，逐條。不是行銷目標。沒有就空陣列"],
  "kol": ["合作的KOL／藝人／代言人姓名。若簡報列的是整個檔期的人選池，全部照列即可"],
  "kol_content": ["簡報有描述的 KOL 演繹內容：穿搭、場景、活動、拍攝設定"],
  "primary_kol": "留空字串——這一欄由人工在「更正」中指定主打代言人，不要自己判斷或推測",
  "activation": "有無實體活動／體驗裝置／快閃店？寫出名稱、內容、地點、時間",
  "consumer_info": "消費者真正需要的資訊：售價、開賣日、販售通路、活動辦法。簡報沒寫就留空",
  "target_audience": "目標受眾描述",
  "mandatory_terms": ["必須出現的字詞，例如品牌 hashtag、產品型號"],
  "forbidden_terms": ["明確禁止使用的字詞（沒有就空陣列）"],
  "internal_only": {{
    "objectives": ["行銷目標、KPI、聲量／業績目標"],
    "media_plan": ["版位、媒體通路、OOH 點位、投放形式"],
    "schedule_phases": "檔期分期（Pre-heat／Launch／Sustain 等）與各期日期",
    "budget": "預算",
    "deliverables": ["交付項目，例如 廣編1則、IG Post 1則、Reels 1支"],
    "other": ["其他只有代理商與品牌需要知道的事"]
  }},
  "uncertain": ["你不確定或簡報中語意模糊的地方，逐條列出供人工確認"]
}}

再次強調：internal_only 底下的東西**不會**被寫進廣編稿。
請確實把投放計畫、檔期分期、KPI、預算放進去，不要留在上面的欄位裡。

【簡報文字】
{text}"""


def extract_text(path: str | Path) -> tuple[str, int]:
    """Return (all text, slide count) from a .pptx."""
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks: list[str] = []
    slide_count = 0

    for i, slide in enumerate(prs.slides):
        slide_count += 1
        parts: list[str] = []

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
                    walk(shape.shapes)
                    continue
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        parts.append(t)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))

        walk(slide.shapes)

        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                parts.append(f"（備註）{note}")

        if parts:
            chunks.append(f"--- 第 {i + 1} 張 ---\n" + "\n".join(parts))

    return "\n\n".join(chunks), slide_count


def _box(shape) -> tuple[int, int, int, int] | None:
    """(left, top, right, bottom) in EMU, or None if the shape is unpositioned."""
    left, top = shape.left, shape.top
    width, height = shape.width, shape.height
    if None in (left, top, width, height):
        return None
    return (left, top, left + width, top + height)


def _gap(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    """Shortest distance between two rectangles; 0 when they overlap.

    This is the geometric half of figure/caption pairing as the document-layout
    literature does it. It is a spatial guess and nothing more: .pptx records
    where each shape sits, never which caption belongs to which picture.
    """
    dx = max(0, a[0] - b[2], b[0] - a[2])
    dy = max(0, a[1] - b[3], b[1] - a[3])
    return (dx * dx + dy * dy) ** 0.5


def extract_images(path: str | Path, max_nearby_chars: int = 300) -> list[dict]:
    """Collect every embedded picture, with the context needed to judge it later.

    Returns one dict per picture: slide number, geometry, the raw bytes, the
    slide heading, and the nearest text on that slide. The nearby text is a
    proximity guess, so it is labelled as a hint everywhere it surfaces rather
    than presented as the picture's caption.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    out: list[dict] = []

    for i, slide in enumerate(prs.slides):
        pictures: list = []
        texts: list[tuple[tuple[int, int, int, int], str]] = []

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shape, "shapes"):
                    walk(shape.shapes)
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures.append(shape)
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    box = _box(shape)
                    if text and box:
                        texts.append((box, text))

        walk(slide.shapes)

        heading = ""
        try:
            if slide.shapes.title is not None:
                heading = (slide.shapes.title.text or "").strip()
        except (AttributeError, ValueError):
            heading = ""

        for shape in pictures:
            # A linked (not embedded) picture has no blob; skip rather than
            # crash the whole upload over one broken reference.
            try:
                image = shape.image
                blob, ext = image.blob, (image.ext or "").lower()
            except (AttributeError, ValueError, KeyError):
                continue

            box = _box(shape)
            nearby = ""
            if box and texts:
                nearest = min(texts, key=lambda t: _gap(box, t[0]))
                nearby = nearest[1][:max_nearby_chars]

            out.append({
                "slide_index": i + 1,
                "left": shape.left, "top": shape.top,
                "width": shape.width, "height": shape.height,
                "blob": blob,
                "ext": ext,
                "slide_heading": heading[:200],
                "nearby_text": nearby,
            })

    return out


def extract_facts(raw_text: str, max_chars: int = 40000) -> dict:
    """Ask the model to normalise deck text into the fact schema."""
    if not raw_text.strip():
        return {}
    text = raw_text[:max_chars]
    return llm.complete_json(
        instructions=INSTRUCTIONS,
        user_input=TASK.format(text=text),
        timeout=timeouts.extract(),
    )


def brief_query_text(facts: dict) -> str:
    """Condense the facts into one line for topical exemplar retrieval."""
    parts: list[str] = []
    for key in ("brand", "campaign", "slogan", "campaign_context"):
        value = facts.get(key)
        if isinstance(value, str) and value.strip():
            parts.append(value.strip())
    for key in ("product", "key_message", "kol"):
        value = facts.get(key)
        if isinstance(value, list):
            parts.extend(str(v) for v in value[:5] if str(v).strip())
    return " ".join(parts)[:1500]
